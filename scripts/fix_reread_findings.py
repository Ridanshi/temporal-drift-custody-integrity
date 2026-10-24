from docx import Document

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)


def replace_span(p, old, new):
    full = p.text
    assert full.count(old) == 1, f"expected exactly 1 match, got {full.count(old)}: {old!r}"
    idx = full.index(old)
    end_idx = idx + len(old)

    pos = 0
    start_run = end_run = None
    for i, r in enumerate(p.runs):
        run_start = pos
        run_end = pos + len(r.text)
        if start_run is None and run_start <= idx < run_end:
            start_run = i
            start_local = idx - run_start
        if run_start < end_idx <= run_end:
            end_run = i
            end_local = end_idx - run_start
        pos = run_end
    assert start_run is not None and end_run is not None

    if start_run == end_run:
        r = p.runs[start_run]
        r.text = r.text[:start_local] + new + r.text[end_local:]
    else:
        p.runs[start_run].text = p.runs[start_run].text[:start_local] + new
        for k in range(start_run + 1, end_run):
            p.runs[k].text = ""
        p.runs[end_run].text = p.runs[end_run].text[end_local:]


# 1. Abstract: 0.93 -> 0.92 YouTube, matching Table 5 / Section 7.2.4 body text
replace_span(
    doc.paragraphs[3],
    "the full system's mean held-out AUC (0.90 WhatsApp, 0.93 YouTube) exceeds the baseline's",
    "the full system's mean held-out AUC (0.90 WhatsApp, 0.92 YouTube) exceeds the baseline's",
)

# 2. Limitations: stale pre-expansion Table 3 description (8 triplets/4 devices) ->
# matches the actual current Table 3 (86 triplets/15 devices), consistent with
# Section 7.2's own caveat paragraph.
replace_span(
    doc.paragraphs[183],
    "Section 7.2's real-data results (Table 3) use only 8 genuine triplets (4 devices, "
    "2 scenes) and originally reused",
    "Section 7.2's real-data results (Table 3) use all 86 available genuine triplets "
    "(15 devices) and reuse",
)

doc.save(PAPER)
print("Fixed abstract AUC typo and stale Limitations Table 3 description.")
