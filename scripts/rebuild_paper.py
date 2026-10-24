import copy, sys, shutil
sys.path.insert(0, "scripts")
from docx import Document
from docx.text.paragraph import Paragraph
import build_lib as BL
from build_lib import B, W
import paper_content_a as A, paper_content_a2 as A2, paper_content_a3 as A3
import paper_content_b1 as B1, paper_content_b2 as B2, paper_content_b3 as B3
import paper_content_c1 as C1, paper_content_c2 as C2, paper_content_c3 as C3
import paper_content_e1 as E1, paper_content_e2 as E2, paper_content_f1 as F1, paper_content_f2 as F2
import paper_content_d1 as D1, paper_content_d2 as D2, paper_content_d3 as D3, paper_content_d4 as D4

SRC = "paper/Temporal-Drift-Custody-Integrity_ComputersSecurity_pre-lineno-backup.docx"
OUT = "paper/Temporal-Drift-Custody-Integrity_ComputersSecurity.docx"
FIG = "paper/new_figs/"
doc = Document(SRC)
for sec in doc.sections:
    for e in sec._sectPr.findall(W + "lnNumType"): sec._sectPr.remove(e)
for r in list(doc.paragraphs[1].runs):
    if r.font.superscript: r._r.getparent().remove(r._r)
for idx in (4, 3, 2):
    if "[" in doc.paragraphs[idx].text:
        el_ = doc.paragraphs[idx]._p; el_.getparent().remove(el_)
ps = doc.paragraphs
def find(prefix, style=None):
    m = [p for p in doc.paragraphs if p.text.strip().startswith(prefix) and (style is None or p.style.name == style)]
    assert m, prefix
    return m[0]
tpl = {
    "body": copy.deepcopy(find("This snapshot framing")._p), "li": copy.deepcopy(find("Tier 1, naive")._p),
    "h1": copy.deepcopy(find("5. Data", "Heading 1")._p), "h2": copy.deepcopy(find("7.1 Calibrated", "Heading 2")._p),
    "tcap": copy.deepcopy(find("Table 22.")._p), "fcap": copy.deepcopy(find("Figure 13.")._p),
    "table": copy.deepcopy([t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Detector"][0]._tbl),
}
def settext(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
def sub(p, a, b):
    for r in p.runs:
        if a in r.text: r.text = r.text.replace(a, b); return
    raise AssertionError(a)
# KEYWORDS_FIX
kw = find("Keywords:")
kw.runs[-1].text = "Digital forensics; Chain of custody; Video authenticity; Anomaly detection; Anti-forensics; Adversarial evaluation"
# ---- earlier sections ----
settext(find("Forensic authenticity checks"), A.ABSTRACT)
settext(find("We reframe chain-of-custody"), A2.INTRO_FRAME)
settext(find("Section 3 formalizes"), A2.INTRO_PLAN)
cp = find("Contributions.")
cp.runs[0].text = "Contributions. "; cp.runs[1].text = A2.CONTRIB
for r in cp.runs[2:]: r.text = ""
settext(find("We evaluate against three attacker tiers"), A3.THREAT_INTRO)
for pre, txt in zip(("Tier 1, naive", "Tier 2, heuristic", "Tier 3, adaptive"), A3.TIERS): settext(find(pre), txt)
settext(find("A detector that only handles Tier 1"), A3.THREAT_END)
sub(find("Every scoring primitive"), "in Section 7.3,", "in Section 7.4,")
sub(find("The preference of Section 1"), "Sections 7.3 and 8.3 revisit this trade-off empirically.", "Section 7.4 revisits this trade-off empirically.")
pp = find("Positioning.")
print([ (r.text[:30]) for r in pp.runs])
full = pp.text
i = full.index("The contribution is the synthesis")
new_tail = "The contribution is the synthesis (per-transform drift baselines, multi-hop graph propagation, and a two-sided co-evasion test) and its evaluation on real forged files against attackers that adapt to the detector."
pos = 0
for r in pp.runs:
    L = len(r.text)
    if pos + L <= i: pos += L; continue
    if pos <= i: r.text = r.text[: i - pos] + new_tail; pos += L; continue
    r.text = ""; pos += L
settext(find("Thresholds are calibrated as a percentile"), B1.DRIFT_P)
fe = find("Intrinsic features.")
settext(fe, B1.FEAT_P)
sub_el = fe._p
settext(find("A single edge"), find("A single edge").text + B2.PATH_EXTRA)
settext(find("This is the paper"), B1.COEV_P)
settext(find("tau_high and tau_low are calibrated"), B1.COEV_P2)
settext(find("The full system combines all three"), B2.AGG_P)
settext(find("Following Xiang"), B2.BASE_P)
dp = find("A larger literature addresses")
settext(dp, dp.text + " Public manipulation corpora such as DFD [43] and Celeb-DF [44] provide manipulated videos but no custody history or transfer chain, so we forge from VISION recordings instead.")
r7 = find("[7] Q. McNemar")
settext(r7, "[7] B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. New York, NY, USA: Chapman & Hall, 1993.")
# ---- remove old sections 5..9 ----
start = find("5. Data", "Heading 1")._p
end = find("CRediT authorship", "Heading 1")._p
el = start; rm = []
while el is not end:
    rm.append(el); el = el.getnext()
for e in rm: e.getparent().remove(e)
settext(find("The source code, unit tests"), C3.AVAIL)
# ---- Table 2 after the feature paragraph ----
bl = B(doc, fe._p.getnext(), tpl)
t = bl.table(*D1.T_FEAT, widths=[1100, 3000, 5260], font=8.5)
for r in t.rows:
    for c in r.cells: c.paragraphs[0].alignment = 0
# ---- new sections 5-9 ----
b = B(doc, end, tpl)
def left(t):
    for r in t.rows:
        for c in r.cells: c.paragraphs[0].alignment = 0
b.heading("5. Data and Forgery Pipeline", 1)
b.para(B2.DATA_1); b.para(B2.DATA_2)
left(b.table(*D1.T_EDIT, widths=[1500, 7860], font=9))
for x in (B3.DATA_3, B3.DATA_4, B3.DATA_5, B3.DATA_6): b.para(x)
b.heading("6. Evaluation Protocol", 1)
for x in C1.PROT: b.para(x)
b.heading("7. Results", 1)
b.heading("7.1 Pipeline fidelity and control", 2); b.para(C1.R71)
b.heading("7.2 Single-hop detection", 2); b.para(C1.R72[0])
b.table(*D2.T_AUC, widths=[1900, 1500, 1500, 1500, 1500, 1460], font=8)
b.figure(FIG + "figA_auc.png", 5.3, D4.CAPS["figA"])
b.para(C1.R72[1])
b.table(*D1.T_TPR, widths=[2200, 1600, 1860, 1850, 1850], font=9)
b.figure(FIG + "figB_roc.png", 5.4, D4.CAPS["figB"])
b.heading("7.3 Effect of the tamper edit", 2); b.para(C1.R73)
b.table(*D2.T_EDITAUC, widths=[1500, 1300, 1150, 1150, 1450, 1350, 1460], font=8.5)
b.figure(FIG + "figC_edits.png", 5.4, D4.CAPS["figC"])
b.heading("7.4 Paired comparisons and learned detectors", 2); b.para(C1.R74)
b.table(*D3.T_PAIR, widths=[2100, 1800, 1800, 1830, 1830], font=8)
b.heading("7.5 What carries detection", 2); b.para(C1.R75)
b.table(*D3.T_ABL, widths=[3200, 1500, 1500, 1580, 1580], font=9)
b.heading("7.6 Weaker edits", 2); b.para(C2.R76)
b.table(*D3.T_STR, widths=[2400, 2300, 2300, 2360], font=9)
b.figure(FIG + "figE_strength.png", 3.9, D4.CAPS["figE"])
b.heading("7.7 The adaptive attacker", 2); b.para(C2.R77[0])
b.figure(FIG + "figF_search.png", 6.0, D4.CAPS["figF"])
b.para(C2.R77[1])
b.heading("7.8 False transfer-type declaration", 2); b.para(C2.R78)
b.table(*D4.T_FD, widths=[2700, 1100, 1200, 1300, 1200, 1860], font=8.5)
b.heading("7.9 Multi-hop chains", 2); b.para(C2.R79[0])
b.table(*D4.T_CH, widths=[1900, 1500, 1300, 1560, 1560, 1540], font=8.5)
b.table(*D4.T_CHAUC, widths=[2200, 1190, 1190, 1190, 1200, 1195, 1195], font=8.5)
b.para(C2.R79[1])
b.figure(FIG + "figD_chains.png", 6.0, D4.CAPS["figD"])
b.heading(E1.H710, 2); b.para(E1.R710[0])
b.table(*E2.T_FUS, widths=[2300, 1400, 1400, 1420, 1420, 1420], font=7.5)
b.para(E1.R710[1])
b.heading(E1.H711, 2); b.para(E1.R711)
b.table(*E2.T_HDR, widths=[2400, 2300, 2300, 2360], font=9)
b.heading(E1.H712, 2); b.para(E1.R712)
b.table(*E2.T_X, widths=[2200, 1500, 1500, 1700, 2460], font=8.5)
b.heading(E1.H713, 2); b.para(E1.R713[0])
b.figure(FIG + "figG_budget.png", 3.9, E1.CAP_G)
b.para(E1.R713[1])
b.heading(F1.H714, 2); b.para(F1.R714[0]); b.para(F1.R714[1])
b.table(*F2.T_GB, widths=[2100, 1210, 1210, 1210, 1210, 1210, 1210], font=7)
b.table(*F2.T_LO, widths=[2400, 2300, 2300, 2360], font=9)
b.para(F1.R714_ATK)
b.heading("8. Discussion", 1)
for h, txt in C3.DISC:
    b.heading(h, 2); b.para(txt)
b.heading("9. Conclusion", 1); b.para(C3.CONCL)
# ---- references ----
last = [p for p in doc.paragraphs if p.text.startswith("[45]")][0]
nxt = last._p.getnext()
for ref in C3.REFS_NEW:
    e = copy.deepcopy(last._p); last._p.addnext(e) if nxt is None else nxt.addprevious(e)
    pr = Paragraph(e, doc._body); settext(pr, ref)
# drop orphaned image parts
import re
xml = doc.element.xml
for rid, rel in list(doc.part.rels.items()):
    if "image" in rel.reltype and rid not in xml:
        del doc.part.rels[rid]
doc.save(OUT)
print("saved")
