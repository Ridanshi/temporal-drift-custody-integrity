import copy

from docx import Document
from docx.table import Table

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)

anchor = None
for p in doc.paragraphs:
    if p.text.strip() == "8. Discussion":
        anchor = p
        break
assert anchor is not None

heading_ref = None
for p in doc.paragraphs:
    if p.text.strip().startswith("7.13 Robustness"):
        heading_ref = p
        break
assert heading_ref is not None

table_caption_ref = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Table 21."):
        table_caption_ref = p
        break
assert table_caption_ref is not None

ref_table = None
for t in doc.tables:
    hdr = [c.text for c in t.rows[0].cells]
    if hdr[0] == "N hops" and len(t.columns) == 5:
        ref_table = t

assert ref_table is not None


def insert_heading_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    p.style = heading_ref.style
    p.add_run(text)
    return p


def insert_body_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    p.add_run(text)
    return p


def insert_table_caption_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    r = p.add_run(text)
    r.italic = True
    r.font.size = table_caption_ref.runs[0].font.size
    return p


def clone_table_before(anchor_para, header, rows):
    new_tbl = copy.deepcopy(ref_table._tbl)
    anchor_para._p.addprevious(new_tbl)
    t = Table(new_tbl, doc)
    n_cols = len(header)
    n_rows_needed = 1 + len(rows)

    while len(t.columns) < n_cols:
        t.add_column(t.columns[0].width if t.columns else None)
    while len(t.columns) > n_cols:
        for row in t.rows:
            row._tr.remove(row._tr.tc_lst[-1])

    while len(t.rows) < n_rows_needed:
        new_row_el = copy.deepcopy(t.rows[-1]._tr)
        t.rows[-1]._tr.addnext(new_row_el)
    while len(t.rows) > n_rows_needed:
        last = t.rows[-1]._tr
        last.getparent().remove(last)

    for c_i, htext in enumerate(header):
        cell = t.rows[0].cells[c_i]
        for p in cell.paragraphs:
            for r in list(p.runs):
                r.text = ""
            if p.runs:
                p.runs[0].text = htext
            else:
                p.add_run(htext)

    for r_i, row_data in enumerate(rows, start=1):
        for c_i, val in enumerate(row_data):
            cell = t.rows[r_i].cells[c_i]
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r.text = ""
                if p.runs:
                    p.runs[0].text = val
                else:
                    p.add_run(val)
    return t


# ============ fix 7.13's closing sentence to point forward to 7.14 ============
target_old_tail = (
    "is a natural response to this finding that this paper does not build or evaluate."
)
target_new_tail = (
    "is a natural response to this finding, and Section 7.14 builds and evaluates it."
)
fixed = False
for p in doc.paragraphs:
    if target_old_tail in p.text:
        full = p.text
        assert full.count(target_old_tail) == 1
        prefix_len = len(full) - len(target_old_tail)
        # find which run boundary prefix_len falls on, splice at run level
        pos = 0
        for i, r in enumerate(p.runs):
            run_start = pos
            run_end = pos + len(r.text)
            if run_start <= prefix_len <= run_end:
                split_in_run = prefix_len - run_start
                r.text = r.text[:split_in_run]
                break
            pos = run_end
        # simplest robust approach: clear all runs after the split point entirely,
        # then append the new tail text to the (now-truncated) split run.
        j = i + 1
        while j < len(p.runs):
            p.runs[j].text = ""
            j += 1
        p.runs[i].text = p.runs[i].text + target_new_tail
        fixed = True
        break
assert fixed, "could not find 7.13 closing sentence to update"

# ============ 7.14 A two-sided elevation test ============
insert_heading_before(anchor, "7.14 A two-sided elevation test")

insert_body_before(
    anchor,
    "Section 7.13 named the natural fix directly: the elevation test only ever checks "
    "whether too MANY hops land above the elevated bar, so an attacker who forces zero "
    "hops above it on every hop evades it by construction, at any chain length. Under a "
    "genuine chain, though, each hop independently has probability (1 - elevated "
    "percentile) of landing above the bar, so seeing zero, or very few, elevated hops "
    "out of many is itself improbable under that same null. "
    "custody.propagate.path_elevation_deficit_pvalue adds exactly this low-tail "
    "binomial check, and flag_path_elevation_two_sided combines both tails, each tested "
    "at alpha/2 (Bonferroni split) so the overall false-positive rate under the null "
    "stays at the same alpha used throughout Sections 7.10 through 7.13. This mirrors "
    "how the co-evasion test already treats both mismatch and suspiciously-perfect "
    "alignment as anomalous, applied here to the elevation test's own count statistic."
)

insert_table_caption_before(
    anchor,
    "Table 22. The Section 7.13 elevation-test-aware attacker (every hop held just "
    "under the elevated bar) against the one-sided elevation test versus the new "
    "two-sided version, combined with the z-test, by chain length.",
)
adaptive = {}
import json as _json
_r = _json.load(open(r"D:\DF paper\scripts\path_two_sided_result.json"))
rows = []
for n in ["1", "3", "5", "7", "9", "11", "15", "30", "50", "75", "100"]:
    wa = _r["whatsapp"]["adaptive_attacker_sweep"][n]
    yt = _r["youtube"]["adaptive_attacker_sweep"][n]
    rows.append([
        n,
        "yes" if wa["old_combined_flagged"] else "no",
        "yes" if wa["new_combined_flagged"] else "no",
        "yes" if yt["old_combined_flagged"] else "no",
        "yes" if yt["new_combined_flagged"] else "no",
    ])
clone_table_before(
    anchor,
    header=[
        "N hops", "WhatsApp: old (1-sided) flags?", "WhatsApp: new (2-sided) flags?",
        "YouTube: old (1-sided) flags?", "YouTube: new (2-sided) flags?",
    ],
    rows=rows,
)

wa_fp15 = _r["whatsapp"]["genuine_fp"]["15"]
wa_fp100 = _r["whatsapp"]["genuine_fp"]["100"]
yt_fp15 = _r["youtube"]["genuine_fp"]["15"]
yt_fp100 = _r["youtube"]["genuine_fp"]["100"]

insert_body_before(
    anchor,
    "The two-sided test closes the specific gap Section 7.13 found, on both edge types, "
    "not just the one (WhatsApp) the old combined detector eventually caught by "
    "accident. The crossover is N = 11 hops on both WhatsApp and YouTube edges, matching "
    "the closed-form expectation exactly: under the null, P(zero of N hops elevated) = "
    f"0.7^N, which first drops below alpha/2 = 0.025 at N = 11 (0.7^10 = 0.0282, 0.7^11 "
    "= 0.0198), independent of which edge type's score distribution is being tested, "
    "because the deficit test operates on the count of elevated hops, not on the "
    "underlying scores' scale or skew. This is a sharp improvement over the one-sided "
    "combined detector, which never caught the YouTube version of this attacker within "
    "the N = 100 hops tested in Section 7.13, and only caught the WhatsApp version at "
    f"N = 75, as an unrelated side effect of the z-test's own low-tail check. The "
    "Bonferroni split does not inflate genuine false alarms: at N = 15 the two-sided "
    f"rate is {wa_fp15['two_sided_rate']:.1%} for WhatsApp and {yt_fp15['two_sided_rate']:.1%} "
    f"for YouTube, versus {wa_fp15['one_sided_rate']:.1%} and {yt_fp15['one_sided_rate']:.1%} "
    f"for the one-sided test; at N = 100 the two-sided rate is {wa_fp100['two_sided_rate']:.1%} "
    f"and {yt_fp100['two_sided_rate']:.1%}, versus {wa_fp100['one_sided_rate']:.1%} and "
    f"{yt_fp100['one_sided_rate']:.1%} one-sided. The two-sided rate is consistently at or "
    "below the one-sided rate at every chain length tested, not above it, since the "
    "stricter alpha/2 bar on the high tail more than offsets the new low-tail checks in "
    "these distributions. custody.pipeline.path_combined_verdict, full_system_verdict, "
    "and drift_only_verdict all now accept an optional path_elevation_deficit flag "
    "alongside path_elevated, contributing a distinct path_elevation_deficit reason, so "
    "a flagged chain's report distinguishes which tail fired rather than collapsing both "
    "into one signal. This does not end the chase: an attacker aware specifically of the "
    "two-sided test's own low-tail threshold could in principle let a small, calibrated "
    "number of hops land just above the bar, exactly enough to avoid both tails at once. "
    "We have not built or tested that next-order attacker; each fix in this chain has so "
    "far produced a new, narrower attacker rather than a final one, which appears to be "
    "an inherent property of this kind of adversarial testing rather than a defect "
    "specific to any one signal tested in this paper."
)

doc.save(PAPER)
print("Inserted section 7.14 and updated the 7.13 forward reference.")
