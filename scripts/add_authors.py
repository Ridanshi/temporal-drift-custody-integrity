from docx import Document

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)

p = doc.paragraphs[1]
assert "Working paper draft" in p.text

# clear existing runs
for r in list(p.runs):
    r.text = ""

p.runs[0].text = "Ridanshi Agarwal"
r_sup1 = p.add_run("1")
r_sup1.font.superscript = True
p.add_run(", Dr. MADIAJAGAN M")
r_sup2 = p.add_run("2")
r_sup2.font.superscript = True

doc.save(PAPER)
print("Replaced working-draft line with author names.")
print(p.text)
