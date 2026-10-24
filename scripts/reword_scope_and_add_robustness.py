import json
from docx import Document

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)


def replace_span(p, old, new):
    full = p.text
    assert full.count(old) == 1, f"expected exactly 1 match, got {full.count(old)}: {old!r}"
    idx = full.index(old)
    end_idx = idx + len(old)

    pos = 0
    start_run = end_run = None
    for i, r in enumerate(p.runs):
        run_start = pos
        run_end = pos + len(r.text)
        if start_run is None and run_start <= idx < run_end:
            start_run = i
            start_local = idx - run_start
        if run_start < end_idx <= run_end:
            end_run = i
            end_local = end_idx - run_start
        pos = run_end
    assert start_run is not None and end_run is not None

    if start_run == end_run:
        r = p.runs[start_run]
        r.text = r.text[:start_local] + new + r.text[end_local:]
    else:
        p.runs[start_run].text = p.runs[start_run].text[:start_local] + new
        for k in range(start_run + 1, end_run):
            p.runs[k].text = ""
        p.runs[end_run].text = p.runs[end_run].text[end_local:]


# ============ 1. header/draft note ============
replace_span(
    doc.paragraphs[1],
    "scaling to the full VISION corpus remains future work.",
    "scaling to the full VISION corpus is out of scope for this pass.",
)

# ============ 2. Section 7.2's real-data caveat paragraph ============
replace_span(
    doc.paragraphs[85],
    "scaling further toward the fold count Section 7.2.4's power analysis specifies "
    "remains the natural next step before that specific claim, held-out or otherwise, "
    "could support a publication claim (Section 8.4).",
    "scaling further toward the fold count Section 7.2.4's power analysis specifies is "
    "out of scope for this pass; that specific claim, held-out or otherwise, does not "
    "rest on formal significance at the current device count (Section 8.4).",
)

# ============ 3. Conclusion ============
replace_span(
    doc.paragraphs[190],
    "The remaining work is scaling further, beyond fifteen devices to the fold counts "
    "the power analysis specifies for formal significance on the adaptive-tier mean "
    "difference, and a real rather than synthetic tampered-evidence corpus for the "
    "adaptive tier specifically, to move from the proof of concept reported here to a "
    "publication-grade claim.",
    "Scaling further, beyond fifteen devices to the fold counts the power analysis "
    "specifies for formal significance on the adaptive-tier mean difference, and a real "
    "rather than synthetic tampered-evidence corpus for the adaptive tier specifically, "
    "is out of scope for this pass; what is reported here is a proof of concept with an "
    "honestly characterized, not yet closed, significance gap, rather than a "
    "publication-grade claim on that specific point.",
)

# ============ 4. robustness-check paragraph after Section 7.2.4's power-analysis discussion ============
rob = json.load(open(r"D:\DF paper\scripts\robustness_tests_n15_result.json"))
wa, yt = rob["whatsapp"], rob["youtube"]

anchor = doc.paragraphs[94]
assert "We do not read this as a weaker result" in anchor.text

new_p = anchor.insert_paragraph_before("")
new_p.style = anchor.style
new_p.add_run(
    "The paired t-test above assumes normally distributed paired differences, an "
    "assumption the baseline's bimodal per-device behavior already puts in doubt. Two "
    "distribution-free checks on the same 15 paired differences, a Wilcoxon signed-rank "
    "test and an exact sign-flip permutation test enumerating all 2^15 sign patterns "
    "directly rather than relying on any asymptotic approximation, test whether the "
    f"t-test's p-value is an artifact of that assumption: Wilcoxon gives p = {wa['wilcoxon_pvalue']:.3f} "
    f"on WhatsApp and p = {yt['wilcoxon_pvalue']:.3f} on YouTube, and the exact permutation test gives "
    f"p = {wa['exact_permutation_pvalue']:.3f} and p = {yt['exact_permutation_pvalue']:.3f} respectively, "
    "each further from significance than the paired t-test, not closer. The honest "
    "reading is that the t-test's non-significance is not a normality artifact "
    "suppressing a real effect; two assumption-free alternatives agree with it, one of "
    "them more strongly. This is evidence against, not for, the idea that a different "
    "test choice would resolve this at the current device count."
)

doc.save(PAPER)
print("Applied 3 wording fixes and inserted the robustness-check paragraph.")
