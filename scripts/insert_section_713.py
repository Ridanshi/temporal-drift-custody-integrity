import copy
import json

from docx import Document
from docx.table import Table

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)

anchor = None
for p in doc.paragraphs:
    if p.text.strip() == "8. Discussion":
        anchor = p
        break
assert anchor is not None

heading_ref = None
for p in doc.paragraphs:
    if p.text.strip().startswith("7.12 Integrating"):
        heading_ref = p
        break
assert heading_ref is not None

table_caption_ref = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Table 19."):
        table_caption_ref = p
        break
assert table_caption_ref is not None

ref_table = doc.tables[17]  # 5 columns, exact match for the new table's shape


def insert_heading_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    p.style = heading_ref.style
    p.add_run(text)
    return p


def insert_body_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    p.add_run(text)
    return p


def insert_table_caption_before(anchor_para, text):
    p = anchor_para.insert_paragraph_before("")
    r = p.add_run(text)
    r.italic = True
    r.font.size = table_caption_ref.runs[0].font.size
    return p


def clone_table_before(anchor_para, header, rows):
    new_tbl = copy.deepcopy(ref_table._tbl)
    anchor_para._p.addprevious(new_tbl)
    t = Table(new_tbl, doc)
    n_cols = len(header)
    n_rows_needed = 1 + len(rows)

    while len(t.columns) < n_cols:
        t.add_column(t.columns[0].width if t.columns else None)
    while len(t.columns) > n_cols:
        for row in t.rows:
            row._tr.remove(row._tr.tc_lst[-1])

    while len(t.rows) < n_rows_needed:
        new_row_el = copy.deepcopy(t.rows[-1]._tr)
        t.rows[-1]._tr.addnext(new_row_el)
    while len(t.rows) > n_rows_needed:
        last = t.rows[-1]._tr
        last.getparent().remove(last)

    for c_i, htext in enumerate(header):
        cell = t.rows[0].cells[c_i]
        for p in cell.paragraphs:
            for r in list(p.runs):
                r.text = ""
            if p.runs:
                p.runs[0].text = htext
            else:
                p.add_run(htext)

    for r_i, row_data in enumerate(rows, start=1):
        for c_i, val in enumerate(row_data):
            cell = t.rows[r_i].cells[c_i]
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r.text = ""
                if p.runs:
                    p.runs[0].text = val
                else:
                    p.add_run(val)
    return t


pa = json.load(open(r"D:\DF paper\scripts\path_adaptive_result.json"))

# ================= 7.13 Robustness to an elevation-test-aware attacker =================
insert_heading_before(anchor, "7.13 Robustness to an elevation-test-aware attacker")

insert_body_before(
    anchor,
    "Section 7.12's attackers were both built to evade the original z-test, not the "
    "elevation test specifically. An attacker who knows the system now includes an "
    "elevation test has a cheaper option than the maximal creeping attacker from Section "
    "7.10: stay strictly under the elevation test's own elevated bar on every hop, "
    "guaranteeing zero hops are ever counted as elevated, so the binomial test never "
    "fires regardless of chain length. We test exactly this attacker, holding every hop "
    "just under the 70th-percentile elevated bar used throughout Sections 7.11 and 7.12."
)

insert_table_caption_before(
    anchor,
    "Table 21. Elevation-test-aware attacker (every hop held just under the 70th-"
    "percentile elevated bar) versus chain length N: z-score of the cumulative path sum "
    "and whether the combined detector (z-test or elevation test) flags it.",
)
rows = []
for n in ["1", "5", "9", "15", "30", "50", "75", "100"]:
    wa = pa["whatsapp"]["sweep"][n]
    yt = pa["youtube"]["sweep"][n]
    rows.append([
        n,
        f"{wa['attack_z']:.2f}",
        "yes" if wa["combined_flagged"] else "no",
        f"{yt['attack_z']:.2f}",
        "yes" if yt["combined_flagged"] else "no",
    ])
clone_table_before(
    anchor,
    header=["N hops", "WhatsApp: z", "WhatsApp: flagged?", "YouTube: z", "YouTube: flagged?"],
    rows=rows,
)

insert_body_before(
    anchor,
    "The elevation test never fires against this attacker on either edge type, at any "
    "chain length tested, by construction, exactly as designed. What happens instead is "
    "unexpected and worth explaining rather than glossing over: both transform types' "
    "genuine score distributions are heavily right-skewed, so the elevated bar itself "
    "sits below the population mean (WhatsApp: bar -1.81 versus mean -0.56; YouTube: bar "
    "-3.95 versus mean -3.56). Holding every hop at a value below the elevated bar is "
    "therefore holding every hop below the population mean too, and a chain that never "
    "benefits from the occasional high-value genuine draw a real bootstrap chain would "
    "eventually include drifts increasingly below what a genuine chain's cumulative sum "
    "would typically reach. This eventually trips the z-test's low-tail check, the same "
    "two-sided |z| > 3 condition already in place, not a new mechanism, but only for "
    "WhatsApp-type edges, crossing at N = 75. YouTube-type edges are not caught within "
    "any chain length tested up to N = 100; the same calibration asymmetry named in "
    "Sections 7.9 and 7.10 shows up here in a third, independent form. Genuine "
    "false-positive rate stays controlled throughout, roughly 2% to 8%, consistent with "
    "every other table in this section. The honest reading is not that the elevation "
    "test failed, it worked exactly as designed and never fires against an attacker who "
    "specifically avoids its trigger condition, but that an attacker aware of the "
    "elevation test's specific mechanism can evade the combined detector entirely for "
    "realistic chain lengths on YouTube edges, and for a substantial number of hops on "
    "WhatsApp edges before an unrelated side effect eventually catches them. A fully "
    "two-sided elevation test, one that also treats an unnaturally low elevated-hop "
    "count as suspicious, mirroring how the co-evasion test treats both mismatch and "
    "suspiciously perfect alignment as anomalous, is a natural response to this finding "
    "that this paper does not build or evaluate."
)

doc.save(PAPER)
print("Inserted section 7.13.")
