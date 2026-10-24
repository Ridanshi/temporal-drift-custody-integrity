import copy
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class B:
    def __init__(self, doc, anchor_el, tpl):
        self.doc, self.anchor, self.t = doc, anchor_el, tpl

    def _clone(self, key):
        el = copy.deepcopy(self.t[key])
        self.anchor.addprevious(el)
        return el

    def para(self, text, key="body"):
        el = self._clone(key)
        p = Paragraph(el, self.doc._body)
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        prev = el.getprevious()
        if key == "body" and prev is not None and prev.tag == W + "tbl":
            p.paragraph_format.space_before = Pt(10)
        return p

    def heading(self, text, level):
        return self.para(text, f"h{level}")

    def bullet(self, text):
        return self.para(text, "li")

    def caption_table(self, text):
        return self.para(text, "tcap")

    def figure(self, path, width, caption):
        el = self._clone("fcap")
        p = Paragraph(el, self.doc._body)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.add_run().add_picture(path, width=Inches(width))
        p.paragraph_format.keep_with_next = True
        cap = self._clone("fcap")
        cp = Paragraph(cap, self.doc._body)
        lab, rest = caption.split(". ", 1)
        cp.runs[0].text = lab + ". "
        cp.runs[1].text = rest
        return p

    def table(self, caption, header, rows, widths=None, font=None):
        _cp = self.caption_table(caption); _cp.paragraph_format.keep_with_next = True
        _pv = _cp._p.getprevious()
        if _pv is not None and _pv.tag == W + "tbl":
            _cp.paragraph_format.space_before = Pt(12)
        el = copy.deepcopy(self.t["table"])
        self.anchor.addprevious(el)
        t = Table(el, self.doc._body)
        n = len(header)
        grid = el.find(W + "tblGrid")
        while len(grid) < n:
            grid.append(copy.deepcopy(grid[-1]))
            for tr in t._tbl.findall(W + "tr"):
                tr.append(copy.deepcopy(tr.findall(W + "tc")[-1]))
        while len(grid) > n:
            grid.remove(grid[-1])
            for tr in t._tbl.findall(W + "tr"):
                tr.remove(tr.findall(W + "tc")[-1])
        trs = t._tbl.findall(W + "tr")
        while len(trs) < 1 + len(rows):
            new = copy.deepcopy(trs[-1]); trs[-1].addnext(new); trs = t._tbl.findall(W + "tr")
        while len(trs) > 1 + len(rows):
            trs[-1].getparent().remove(trs[-1]); trs = t._tbl.findall(W + "tr")
        total = 9360
        w = widths or ([2400] + [int((total - 2400) / (n - 1))] * (n - 1))
        for g, x in zip(grid, w):
            g.set(W + "w", str(x))
        tw = el.find(W + "tblPr").find(W + "tblW")
        tw.set(W + "w", str(sum(w))); tw.set(W + "type", "dxa")
        for ri, data in enumerate([header] + rows):
            tcs = trs[ri].findall(W + "tc")
            for ci, val in enumerate(data):
                tc = tcs[ci]
                tcw = tc.find(W + "tcPr").find(W + "tcW") if tc.find(W + "tcPr") is not None else None
                if tcw is not None:
                    tcw.set(W + "w", str(w[ci])); tcw.set(W + "type", "dxa")
                p = Paragraph(tc.findall(W + "p")[0], t._parent)
                if not p.runs:
                    p.add_run("")
                p.runs[0].text = val
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                if font:
                    p.runs[0].font.size = Pt(font)
                if ci == 0 and len(set(len(r[0]) for r in rows)) >= 1 and header[0] in ("Group", "Edit", "Method", "Case", "Chains", "Statistic", "Comparison", "Feature set"):
                    p.alignment = 0
        for ri in range(len(trs) - 1):
            for pp_ in trs[ri].iter(W + "p"):
                Paragraph(pp_, t._parent).paragraph_format.keep_with_next = True
        for tr in trs:
            trPr = tr.find(W + "trPr")
            if trPr is None:
                from docx.oxml import OxmlElement
                trPr = OxmlElement("w:trPr"); tr.insert(0, trPr)
            if trPr.find(W + "cantSplit") is None:
                from docx.oxml import OxmlElement
                trPr.append(OxmlElement("w:cantSplit"))
        return t
