from docx import Document

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)

target = None
for p in doc.paragraphs:
    if "Every path-anomaly defense" in p.text:
        target = p
        break
assert target is not None

addition = (
    " Section 7.14 closes the specific gap Section 7.13 found by adding a low-tail "
    "check to the elevation test, catching that attacker at 11 hops on both edge types "
    "rather than 75 hops or never, without inflating the genuine false-positive rate; it "
    "does not close the chase itself, an attacker aware of the two-sided test's own "
    "low-tail threshold specifically could in principle still be constructed, and has "
    "not been tested here."
)

last_run = target.runs[-1]
new_run = target.add_run(addition)
new_run.font.name = last_run.font.name
new_run.font.size = last_run.font.size
new_run.bold = last_run.bold
new_run.italic = last_run.italic

doc.save(PAPER)
print("Updated limitations paragraph for 7.14.")
