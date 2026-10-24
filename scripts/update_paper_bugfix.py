from docx import Document

PAPER = r"D:\DF paper\paper\Temporal-Drift-Custody-Integrity.docx"
doc = Document(PAPER)
P = doc.paragraphs


def set_single_run(idx, text):
    p = P[idx]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_row_cells(row, values):
    for ci, val in enumerate(values, start=1):
        cell = row.cells[ci]
        for p in cell.paragraphs:
            for r in list(p.runs):
                r.text = ""
            if p.runs:
                p.runs[0].text = val
            else:
                p.add_run(val)


# --- Table 3 (doc.tables[2]): fix coevasion_only cells only ---
t3 = doc.tables[2]
for row in t3.rows[1:]:
    if row.cells[0].text.strip() == "Co-evasion-only (ablation)":
        set_row_cells(row, ["0.79 / 0.79", "0.94 / 0.94", "0.45 / 0.45"])

# --- Table 4 (doc.tables[3]): McNemar, corrected ---
t4 = doc.tables[3]
table4_data = {
    "WhatsApp / naive": ["11", "8", "0.21", "0.646"],
    "WhatsApp / heuristic": ["31", "8", "12.41", "4.3e-04"],
    "WhatsApp / adaptive": ["26", "8", "8.50", "3.6e-03"],
    "YouTube / naive": ["14", "9", "0.70", "0.404"],
    "YouTube / heuristic": ["34", "9", "13.40", "2.5e-04"],
    "YouTube / adaptive": ["30", "9", "10.26", "1.4e-03"],
}
for row in t4.rows[1:]:
    key = row.cells[0].text.strip()
    if key in table4_data:
        set_row_cells(row, table4_data[key])
    else:
        print("WARNING unmatched Table 4 row:", key)

# --- Table 7 (doc.tables[6]): DFD, corrected ---
t7 = doc.tables[6]
table7_data = {
    "Snapshot-fusion baseline": ["0.91 / 0.91", "0.94 / 0.97", "0.79 / 0.77", "[0.89, 0.98] / [0.94, 0.99]"],
    "Drift-only (ablation)": ["1.00 / 1.00", "0.97 / 1.00", "0.80 / 0.87", "[0.94, 0.99] / [0.99, 1.00]"],
    "Co-evasion-only (ablation)": ["0.79 / 0.79", "0.75 / 0.75", "0.47 / 0.47", "[0.65, 0.84] / [0.66, 0.84]"],
    "Full system": ["1.00 / 1.00", "0.94 / 1.00", "0.77 / 0.80", "[0.88, 0.99] / [0.99, 1.00]"],
}
for row in t7.rows[1:]:
    method = row.cells[0].text.strip()
    if method in table7_data:
        set_row_cells(row, table7_data[method])
    else:
        print("WARNING unmatched Table 7 row:", method)

# --- Table 8 (doc.tables[7]): Celeb-DF, corrected ---
t8 = doc.tables[7]
table8_data = {
    "Snapshot-fusion baseline": ["0.86 / 0.86", "0.95 / 0.97", "0.83 / 0.81", "[0.92, 0.98] / [0.95, 0.99]"],
    "Drift-only (ablation)": ["1.00 / 1.00", "1.00 / 1.00", "0.90 / 0.90", "[1.00, 1.00] / [1.00, 1.00]"],
    "Co-evasion-only (ablation)": ["0.94 / 0.94", "0.50 / 0.50", "0.11 / 0.11", "[0.46, 0.56] / [0.46, 0.56]"],
    "Full system": ["1.00 / 1.00", "1.00 / 1.00", "0.84 / 0.84", "[1.00, 1.00] / [1.00, 1.00]"],
}
for row in t8.rows[1:]:
    method = row.cells[0].text.strip()
    if method in table8_data:
        set_row_cells(row, table8_data[method])
    else:
        print("WARNING unmatched Table 8 row:", method)

# --- Table 12 (doc.tables[11]): worked case study, corrected ---
t12 = doc.tables[11]
table12_data = {
    "Native intrinsic feature vector": ["[1.121, 0.298, 710.0, 91.0]"],
    "WhatsApp intrinsic feature vector": ["[0.848, 0.385, 710.0, 91.0]"],
    "Delta (WhatsApp minus native)": ["[-0.273, 0.087, 0.0, 0.0]"],
    "Drift score (negative log-likelihood)": ["-1.994"],
    "Drift threshold (90th percentile, genuine)": ["3.029"],
    "Co-evasion Mahalanobis distance": ["1.942"],
    "Co-evasion band [tau_low, tau_high]": ["[0.591, 3.296]"],
    "Co-evasion flag": ["none (inside normal band)"],
    "Full-system verdict": ["not flagged, reasons: []"],
}
for row in t12.rows[1:]:
    key = row.cells[0].text.strip()
    if key in table12_data:
        set_row_cells(row, table12_data[key])
    else:
        print("WARNING unmatched Table 12 row:", key)

doc.save(PAPER)
print("Saved all table corrections.")
